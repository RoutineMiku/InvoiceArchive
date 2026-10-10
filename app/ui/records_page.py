# -*- coding: utf-8 -*-
"""归档记录页：查看、检索、修改（支付时间 / 费用类型 / 备注）、删除、导出。"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PySide6.QtCore import QDate, Qt, Signal
from PySide6.QtGui import QBrush, QColor, QFont
from PySide6.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QDateEdit,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QStyledItemDelegate,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from ..config import (
    AUTO_COLUMNS,
    COL_AMOUNT,
    COL_BEARER,
    COL_EXPENSE,
    COL_PAY_DATE,
    COL_SEQ,
    DEFAULT_EXPENSE_TYPES,
    EDITABLE_COLUMNS,
    RECORD_COLUMNS,
    RECORD_WIDTHS,
)
from ..excel_store import ExcelError
from .bearers_page import open_in_explorer
from .theme import DANGER, TEXT_MUTED, WARNING
from .widgets import Card


class ComboDelegate(QStyledItemDelegate):
    """费用类型列：下拉选择，同时允许自由输入。"""

    def __init__(self, items: list[str], parent=None):
        super().__init__(parent)
        self.items = items

    def createEditor(self, parent, option, index):  # noqa: N802
        editor = QComboBox(parent)
        editor.setEditable(True)
        editor.addItems(self.items)
        return editor

    def setEditorData(self, editor, index):  # noqa: N802
        editor.setCurrentText(index.data() or "")

    def setModelData(self, editor, model, index):  # noqa: N802
        model.setData(index, editor.currentText(), Qt.EditRole)


class PayDateDialog(QDialog):
    """批量设置支付时间。"""

    def __init__(self, parent=None, count: int = 1):
        super().__init__(parent)
        self.setWindowTitle("设置支付时间")
        self.setMinimumWidth(320)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(18, 18, 18, 14)
        lay.setSpacing(10)
        lay.addWidget(QLabel(f"将选中 {count} 条记录的「支付时间」设置为："))
        self.date = QDateEdit(QDate.currentDate())
        self.date.setCalendarPopup(True)
        self.date.setDisplayFormat("yyyy-MM-dd")
        lay.addWidget(self.date)
        self.chk_clear = QCheckBox("清空支付时间（置为空白）")
        lay.addWidget(self.chk_clear)
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.button(QDialogButtonBox.Ok).setText("确定")
        btns.button(QDialogButtonBox.Cancel).setText("取消")
        btns.accepted.connect(self.accept)
        btns.rejected.connect(self.reject)
        lay.addWidget(btns)

    def value(self) -> str:
        if self.chk_clear.isChecked():
            return ""
        return self.date.date().toString("yyyy-MM-dd")


class RecordsPage(QWidget):
    data_changed = Signal()

    def __init__(self, store, parent: QWidget | None = None):
        super().__init__(parent)
        self.store = store
        self._rows: list[dict] = []
        self._loading = False
        self._build_ui()
        self.reload()

    # ------------------------------------------------------------ UI
    def _build_ui(self) -> None:
        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(12)

        # 统计条
        stat_card = Card()
        stat_row = QHBoxLayout()
        stat_row.setSpacing(28)
        self.lb_count = self._make_stat(stat_row, "归档条数")
        self.lb_amount = self._make_stat(stat_row, "价税合计总额（元）")
        self.lb_paid = self._make_stat(stat_row, "已登记支付时间")
        self.lb_unpaid = self._make_stat(stat_row, "未登记支付时间")
        stat_row.addStretch(1)
        stat_card.add_layout(stat_row)
        root.addWidget(stat_card)

        card = Card("归档明细", "数据保存在 data/发票报销归档.xlsx")

        bar = QHBoxLayout()
        bar.setSpacing(8)
        self.ed_search = QLineEdit()
        self.ed_search.setPlaceholderText("搜索：费用承担人 / 发票号码 / 销售方 / 费用类型 / 备注…")
        self.ed_search.setClearButtonEnabled(True)
        self.ed_search.textChanged.connect(self.reload)
        bar.addWidget(self.ed_search, 1)

        self.btn_refresh = QPushButton("刷新")
        self.btn_refresh.clicked.connect(self.reload)
        self.btn_pay = QPushButton("设置支付时间")
        self.btn_pay.setObjectName("Primary")
        self.btn_pay.clicked.connect(self.set_pay_date)
        self.btn_del = QPushButton("删除选中")
        self.btn_del.setObjectName("Danger")
        self.btn_del.clicked.connect(self.delete_selected)
        self.btn_export = QPushButton("导出为…")
        self.btn_export.clicked.connect(self.export_as)
        self.btn_open = QPushButton("打开 Excel")
        self.btn_open.setObjectName("Ghost")
        self.btn_open.clicked.connect(lambda: open_in_explorer(self.store.record_file))
        for b in (self.btn_refresh, self.btn_pay, self.btn_del, self.btn_export, self.btn_open):
            bar.addWidget(b)
        card.add_layout(bar)

        self.table = QTableWidget(0, len(RECORD_COLUMNS))
        self.table.setHorizontalHeaderLabels(RECORD_COLUMNS)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setVisible(False)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.table.setEditTriggers(
            QAbstractItemView.DoubleClicked
            | QAbstractItemView.EditKeyPressed
            | QAbstractItemView.AnyKeyPressed
        )
        h = self.table.horizontalHeader()
        h.setSectionResizeMode(QHeaderView.Interactive)
        for i, w in enumerate(RECORD_WIDTHS):
            self.table.setColumnWidth(i, w * 8)
        h.setSectionResizeMode(RECORD_COLUMNS.index("销售方公司名称"), QHeaderView.Stretch)
        self.table.itemChanged.connect(self._on_item_changed)
        self.expense_delegate = ComboDelegate(list(DEFAULT_EXPENSE_TYPES), self.table)
        self.table.setItemDelegateForColumn(
            RECORD_COLUMNS.index(COL_EXPENSE), self.expense_delegate)
        card.add_widget(self.table)

        tip = QLabel(
            "可直接双击单元格修改「费用承担人 / 价税合计 / 费用类型 / 支付时间 / 备注」，"
            "修改即时写回 Excel；其余列由系统自动生成。"
        )
        tip.setObjectName("Muted")
        tip.setWordWrap(True)
        card.add_widget(tip)

        root.addWidget(card, 1)

    def _make_stat(self, layout: QHBoxLayout, title: str) -> QLabel:
        box = QVBoxLayout()
        box.setSpacing(0)
        lb = QLabel("0")
        lb.setObjectName("StatValue")
        cap = QLabel(title)
        cap.setObjectName("StatLabel")
        box.addWidget(lb)
        box.addWidget(cap)
        layout.addLayout(box)
        return lb

    # ------------------------------------------------------------ 数据
    def reload(self) -> None:
        keyword = self.ed_search.text().strip().lower()
        try:
            all_rows = self.store.list_records()
        except Exception as exc:  # noqa: BLE001
            QMessageBox.warning(self, "读取失败", str(exc))
            return

        if keyword:
            self._rows = [
                r for r in all_rows
                if keyword in " ".join(str(r.get(c, "")) for c in RECORD_COLUMNS).lower()
            ]
        else:
            self._rows = all_rows

        self._loading = True
        self.table.setRowCount(0)
        for rec in self._rows:
            self._append(rec)
        self._loading = False

        # 统计
        summary = self.store.summarize(all_rows)
        paid = sum(1 for r in all_rows if str(r.get(COL_PAY_DATE, "")).strip())
        self.lb_count.setText(str(summary["count"]))
        self.lb_amount.setText(f"{summary['amount']:,.2f}")
        self.lb_paid.setText(str(paid))
        self.lb_unpaid.setText(str(summary["count"] - paid))

        # 费用类型下拉候选：内置 + 历史出现过的
        used = {r.get(COL_EXPENSE, "") for r in all_rows if r.get(COL_EXPENSE)}
        merged = list(DEFAULT_EXPENSE_TYPES)
        for u in sorted(used):
            if u not in merged:
                merged.append(u)
        self.expense_delegate.items = merged

    def _append(self, rec: dict) -> None:
        r = self.table.rowCount()
        self.table.insertRow(r)
        for c, col in enumerate(RECORD_COLUMNS):
            item = QTableWidgetItem(str(rec.get(col, "")))
            if col not in EDITABLE_COLUMNS:
                item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                item.setForeground(QBrush(QColor(TEXT_MUTED)))
            else:
                item.setToolTip("双击修改")
                item.setFont(QFont("Microsoft YaHei UI", 10))
            if col == COL_AMOUNT:
                item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            if col == COL_SEQ:
                item.setTextAlignment(Qt.AlignCenter)
            self.table.setItem(r, c, item)

    # ------------------------------------------------------------ 编辑回写
    def _on_item_changed(self, item: QTableWidgetItem) -> None:
        if self._loading:
            return
        col_name = RECORD_COLUMNS[item.column()]
        if col_name not in EDITABLE_COLUMNS:
            return
        row = item.row()
        if not (0 <= row < len(self._rows)):
            return
        rec = self._rows[row]
        value = item.text().strip()
        if col_name == COL_AMOUNT and value:
            try:
                value = f"{float(value.replace(',', '')):.2f}"
            except ValueError:
                QMessageBox.warning(self, "格式有误", "「价税合计」必须是数字，例如 1234.56")
                self._loading = True
                item.setText(rec.get(COL_AMOUNT, ""))
                self._loading = False
                return
        try:
            self.store.update_record(rec["_row"], {col_name: value})
        except ExcelError as exc:
            QMessageBox.warning(self, "保存失败", str(exc))
            return
        rec[col_name] = value
        if col_name == COL_AMOUNT:
            self._loading = True
            item.setText(value)
            self._loading = False
        if col_name == COL_PAY_DATE:
            paid = sum(1 for r in self._rows if str(r.get(COL_PAY_DATE, "")).strip())
            self.lb_paid.setText(str(paid))
            self.lb_unpaid.setText(str(len(self._rows) - paid))
        self.data_changed.emit()

    # ------------------------------------------------------------ 批量操作
    def _selected_rows(self) -> list[int]:
        return sorted({i.row() for i in self.table.selectedIndexes()})

    def set_pay_date(self) -> None:
        rows = self._selected_rows()
        if not rows:
            QMessageBox.information(self, "提示", "请先选中要设置支付时间的记录（可多选）。")
            return
        dlg = PayDateDialog(self, len(rows))
        if dlg.exec() != QDialog.Accepted:
            return
        value = dlg.value()
        self._loading = True
        for r in rows:
            rec = self._rows[r]
            try:
                self.store.update_record(rec["_row"], {COL_PAY_DATE: value})
            except ExcelError as exc:
                self._loading = False
                QMessageBox.warning(self, "保存失败", str(exc))
                return
            rec[COL_PAY_DATE] = value
            item = self.table.item(r, RECORD_COLUMNS.index(COL_PAY_DATE))
            if item:
                item.setText(value)
        self._loading = False
        paid = sum(1 for x in self._rows if str(x.get(COL_PAY_DATE, "")).strip())
        self.lb_paid.setText(str(paid))
        self.lb_unpaid.setText(str(len(self._rows) - paid))
        self.data_changed.emit()

    def delete_selected(self) -> None:
        rows = self._selected_rows()
        if not rows:
            QMessageBox.information(self, "提示", "请先选中要删除的记录。")
            return
        ans = QMessageBox.question(
            self, "确认删除",
            f"确定从归档表中删除选中的 {len(rows)} 条记录吗？此操作不可恢复"
            "（系统会在 data/backup 中保留当日备份）。",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No,
        )
        if ans != QMessageBox.Yes:
            return
        excel_rows = [self._rows[r]["_row"] for r in rows]
        try:
            self.store.delete_records(excel_rows)
        except ExcelError as exc:
            QMessageBox.warning(self, "删除失败", str(exc))
            return
        self.reload()
        self.data_changed.emit()

    def export_as(self) -> None:
        """把当前（筛选后的）视图导出为一份新的 xlsx。"""
        default = f"发票报销归档_导出_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        path, _ = QFileDialog.getSaveFileName(
            self, "导出为", str(Path.home() / default), "Excel 工作簿 (*.xlsx)")
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "归档明细"
        ws.append(RECORD_COLUMNS)
        header_fill = PatternFill("solid", fgColor="1F6FEB")
        for c in range(1, len(RECORD_COLUMNS) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for rec in self._rows:
            ws.append([rec.get(c, "") for c in RECORD_COLUMNS])
        for i, w in enumerate(RECORD_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 24
        ws.freeze_panes = "A2"
        try:
            wb.save(path)
        except PermissionError:
            QMessageBox.warning(self, "导出失败", "目标文件正被占用，请关闭后重试。")
            return
        QMessageBox.information(self, "导出成功", f"已导出 {len(self._rows)} 条记录到：\n{path}")
        open_in_explorer(path)
